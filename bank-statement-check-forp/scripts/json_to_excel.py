#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
银行流水核查：解析 LLM 整理的交易明细 JSON，执行四步确定性处理后再生成 Excel，
并将文件直接写盘返回给 LLM（不上传平台、不生成 download_url）。

处理顺序：
  ① 多银行流水格式标准化（无正负时按摘要推断收支：货款/回款→收入，材料款/采购款→支出）
  ② 关联方交易识别
  ③ 异常交易筛查
  ④ 资金闭环核查

输入 tables_json 格式（任选其一）：
  1. {"headers": [...], "rows": [[...], ...]}
  2. [{"交易日期":"...", ...}, ...]
  3. [["交易日期",...], ["2024-01-01",...], ...]

可选参数：
  related_parties / 关联方清单：字符串数组或对象数组
  config / 阈值配置：覆盖默认阈值

依赖库：pip install openpyxl
"""

from __future__ import annotations

import io
import json
import logging
import os
import re
from collections import defaultdict
from copy import deepcopy
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

logger = logging.getLogger(__name__)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None
    Font = None
    PatternFill = None
    Alignment = None
    get_column_letter = None

# ---------------------------------------------------------------------------
# 标准字段与别名
# ---------------------------------------------------------------------------
STANDARD_HEADERS = [
    "交易日期", "交易时间", "摘要", "交易金额",
    "对方账号", "对方户名", "当前账户", "当前户名", "银行",
]

HEADER_ALIASES = {
    "交易日期": ["交易日", "日期", "记账日期", "发生日期", "入账日期", "起息日"],
    "交易时间": ["时间", "记账时间", "发生时间", "交易时刻", "入账时间", "交易时点"],
    "摘要": ["用途", "备注", "说明", "交易摘要", "交易说明", "附言", "业务摘要", "摘要信息", "交易类型"],
    "交易金额": ["金额", "发生额", "交易额", "收入", "支出", "借方金额", "贷方金额", "借方发生额", "贷方发生额"],
    "对方账号": ["对方账户", "对手账号", "收款账号", "付款账号", "对方卡号", "交易对手账号"],
    "对方户名": ["对方名称", "对手户名", "收款人", "付款人", "对方姓名", "对手名称", "收款户名", "付款户名", "交易对手"],
    "当前账户": ["本方账号", "账户", "账号", "卡号", "账户号码", "本方账户", "查询账号", "银行账号"],
    "当前户名": ["本方户名", "账户名称", "户名", "本方名称", "账户名", "本方姓名"],
    "银行": ["开户行", "开户银行", "银行名称", "所属银行", "银行名", "本方银行", "账户银行"],
    "借贷标志": ["借贷", "借贷方向", "收付标志", "借贷标记", "借/贷", "方向"],
}

NORMALIZED_HEADERS = [
    "交易日期", "交易时间", "摘要", "收入", "支出",
    "对方账号", "对方户名", "当前账户", "当前户名", "银行",
]

RELATED_HEADERS = NORMALIZED_HEADERS + ["是否关联方", "风险等级", "关联方类型"]
ANOMALY_HEADERS = RELATED_HEADERS + ["是否异常", "命中原因"]
CHAIN_HEADERS = ["起点账户", "起点户名", "起点银行", "支付日期", "支付金额", "路径对手", "路径摘要", "终点账户", "终点户名", "终点日期", "终点金额", "链条类型", "线索说明"]

DEFAULT_CONFIG = {
    "cash_withdraw_threshold": 500_000,       # R1 大额取现
    "high_freq_amount": 1_000_000,            # R2 高频往来累计金额
    "high_freq_days": 7,                      # R2 窗口天数
    "night_start": "23:00",                   # R3
    "night_end": "05:00",                     # R3
    "sensitive_keywords": ["借款", "往来款", "往来", "暂付", "备用金"],
    "large_payment_threshold": 1_000_000,     # 第四步 / 中风险大额
    "trace_window_days": 7,                   # 第四步追踪窗口
    "related_large_threshold": 1_000_000,     # 关联企业大额往来
}

# 经营性质摘要：R2 排除
BUSINESS_KEYWORDS = [
    "货款", "采购款", "回款", "服务费", "税费", "工资", "社保", "公积金",
    "租金", "电费", "水费", "物业费", "利息", "手续费", "分红",
]

# 非经营往来摘要（R2 纳入）
NON_BUSINESS_KEYWORDS = ["往来款", "借款", "暂付", "备用金", "往来", "拆借", "垫付", "代付"]

# 金额无正负号时，按摘要推断收支方向（先匹配支出，再匹配收入）
# 收入（正）：货款、回款等；支出（负）：材料款、采购款等
EXPENSE_SUMMARY_KEYWORDS = [
    "材料款", "采购款", "材料费", "采购", "付货款", "支付货款",
    "付材料", "材料采购", "货款支付",
]
INCOME_SUMMARY_KEYWORDS = [
    "回款", "货款", "收款", "销售收入", "销售款", "到账款", "货款回笼",
]

# 对手方名称噪声后缀
NAME_SUFFIX_NOISE = [
    r"（本部）", r"\(本部\)", r"-分行", r"—分行", r"－分行",
    r"（总部）", r"\(总部\)", r"（分公司）", r"\(分公司\)",
]
LEGAL_ENTITY_SUFFIXES = [
    "股份有限公司", "有限责任公司", "有限公司", "集团有限公司",
    "股份公司", "集团公司", "集团",
]

FILL_RELATED = "FFF2CC"       # 浅黄：关联方
FILL_HIGH = "FFCDD2"          # 浅红：高风险/异常高
FILL_MEDIUM = "FFE0B2"        # 浅橙：中风险
FILL_HEADER = "D9E2F3"


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------
def _resolve_output_path(file_name: str, context: Dict[str, Any]) -> str:
    """解析 Excel 落盘路径：优先工作目录，保证 LLM/沙箱能直接拿到文件。"""
    workdir = (
        context.get("workdir")
        or context.get("cwd")
        or (context.get("extra") or {}).get("workdir")
        or (context.get("extra") or {}).get("cwd")
        or os.getcwd()
    )
    return os.path.abspath(os.path.join(str(workdir), file_name))


def _check_deps() -> List[str]:
    missing = []
    if Workbook is None:
        missing.append("openpyxl")
    return missing


def _norm_header(name: str) -> str:
    return re.sub(r"\s+", "", str(name or "")).lower()


def _map_header(name: str) -> Optional[str]:
    n = _norm_header(name)
    if not n:
        return None
    for std, aliases in HEADER_ALIASES.items():
        if _norm_header(std) == n:
            return std
        for alias in aliases:
            if _norm_header(alias) == n:
                return std
    return None


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    return "" if s in ("—", "-", "None", "null") else s


def _parse_amount(value: Any) -> Optional[float]:
    s = _cell_str(value)
    if not s:
        return None
    s = s.replace(",", "").replace("，", "").replace("¥", "").replace("￥", "").replace("元", "")
    s = re.sub(r"\s+", "", s)
    # 会计括号负数 (123.45)
    if re.fullmatch(r"\([\d.]+\)", s):
        s = "-" + s[1:-1]
    try:
        return float(s)
    except ValueError:
        return None


def _fmt_amount(value: Optional[float]) -> str:
    if value is None or abs(value) < 1e-12:
        return ""
    return f"{value:.2f}"


def _amount_has_explicit_sign(amount_raw: str) -> bool:
    """金额原文是否已带正负号 / 会计括号负数。"""
    s = _cell_str(amount_raw)
    if not s:
        return False
    s = s.replace(",", "").replace("，", "").replace("¥", "").replace("￥", "").replace("元", "")
    s = re.sub(r"\s+", "", s)
    if re.fullmatch(r"\([\d.]+\)", s):
        return True
    return bool(re.match(r"^[+-]", s))


def _infer_direction_from_summary(summary: str) -> Optional[str]:
    """按摘要推断收支：返回 'income' / 'expense' / None。

    规则：材料款/采购款等 → 支出；货款/回款等 → 收入。先匹配支出关键词。
    """
    text = _cell_str(summary)
    if not text:
        return None
    for kw in EXPENSE_SUMMARY_KEYWORDS:
        if kw in text:
            return "expense"
    for kw in INCOME_SUMMARY_KEYWORDS:
        if kw in text:
            return "income"
    return None


def _parse_date(value: Any) -> Optional[str]:
    """统一为 YYYY-MM-DD。"""
    s = _cell_str(value)
    if not s:
        return None
    s = s.replace("年", "-").replace("月", "-").replace("日", "")
    s = s.replace("/", "-").replace(".", "-").replace("—", "-").replace("–", "-")
    s = re.sub(r"\s+.*$", "", s)  # 去掉时间部分
    s = s.strip("-")

    if re.fullmatch(r"\d{8}", s):
        try:
            return datetime.strptime(s, "%Y%m%d").strftime("%Y-%m-%d")
        except ValueError:
            return None

    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return datetime(y, mo, d).strftime("%Y-%m-%d")
        except ValueError:
            return None

    for fmt in ("%Y-%m-%d", "%Y%m%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _parse_time(value: Any) -> Optional[str]:
    s = _cell_str(value)
    if not s:
        return None
    s = s.replace("：", ":").strip()
    # 从日期时间串中提取
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if m:
        hh, mm, ss = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
        if 0 <= hh <= 23 and 0 <= mm <= 59:
            return f"{hh:02d}:{mm:02d}:{ss:02d}"
    if re.fullmatch(r"\d{4,6}", s):
        if len(s) == 4:
            return f"{s[:2]}:{s[2:]}:00"
        if len(s) == 6:
            return f"{s[:2]}:{s[2:4]}:{s[4:]}"
    return None


def _clean_counterparty_name(name: str) -> str:
    """去首尾空格；剔除噪声后缀；统一简称（去法人后缀）。"""
    s = _cell_str(name)
    if not s:
        return ""
    for pat in NAME_SUFFIX_NOISE:
        s = re.sub(pat, "", s)
    s = re.sub(r"\s+", "", s)
    for suf in LEGAL_ENTITY_SUFFIXES:
        if s.endswith(suf) and len(s) > len(suf) + 1:
            s = s[: -len(suf)]
            break
    return s.strip()


def _merge_config(raw: Any) -> Dict[str, Any]:
    cfg = deepcopy(DEFAULT_CONFIG)
    if not raw:
        return cfg
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError:
            return cfg
    if not isinstance(raw, dict):
        return cfg
    key_aliases = {
        "取现阈值": "cash_withdraw_threshold",
        "高频往来阈值": "high_freq_amount",
        "高频往来天数": "high_freq_days",
        "深夜开始": "night_start",
        "深夜结束": "night_end",
        "敏感关键词": "sensitive_keywords",
        "大额支付阈值": "large_payment_threshold",
        "追踪窗口": "trace_window_days",
        "关联大额阈值": "related_large_threshold",
    }
    for k, v in raw.items():
        key = key_aliases.get(k, k)
        if key in cfg and v is not None:
            if key in (
                "cash_withdraw_threshold",
                "high_freq_amount",
                "high_freq_days",
                "large_payment_threshold",
                "trace_window_days",
                "related_large_threshold",
            ):
                try:
                    cfg[key] = float(v) if "days" not in key else int(v)
                    if "days" in key or key == "trace_window_days" or key == "high_freq_days":
                        cfg[key] = int(float(v))
                except (TypeError, ValueError):
                    continue
            elif key == "sensitive_keywords":
                if isinstance(v, str):
                    cfg[key] = [x.strip() for x in re.split(r"[,，、]", v) if x.strip()]
                elif isinstance(v, list):
                    cfg[key] = [str(x).strip() for x in v if str(x).strip()]
            else:
                cfg[key] = v
    return cfg


def _parse_related_parties(raw: Any) -> List[Dict[str, Any]]:
    """
    解析关联方清单。
    支持：
      ["张三", "辰源贸易"]
      [{"name":"张三","type":"实控人","aliases":["张*"]}, ...]
      JSON 字符串
    """
    if not raw:
        return []
    if isinstance(raw, str):
        raw = raw.strip()
        if not raw:
            return []
        try:
            raw = json.loads(raw)
        except json.JSONDecodeError:
            # 逗号分隔
            raw = [x.strip() for x in re.split(r"[,，、\n]", raw) if x.strip()]

    parties: List[Dict[str, Any]] = []
    if isinstance(raw, dict):
        raw = raw.get("related_parties") or raw.get("parties") or raw.get("list") or []

    if not isinstance(raw, list):
        return []

    for item in raw:
        if isinstance(item, str):
            name = item.strip()
            if name:
                parties.append({
                    "name": name,
                    "clean": _clean_counterparty_name(name),
                    "type": "关联方",
                    "aliases": [],
                })
        elif isinstance(item, dict):
            name = _cell_str(item.get("name") or item.get("名称") or item.get("户名") or "")
            if not name:
                continue
            ptype = _cell_str(item.get("type") or item.get("类型") or item.get("角色") or "关联方")
            aliases_raw = item.get("aliases") or item.get("别名") or []
            if isinstance(aliases_raw, str):
                aliases = [x.strip() for x in re.split(r"[,，、]", aliases_raw) if x.strip()]
            else:
                aliases = [_cell_str(x) for x in aliases_raw if _cell_str(x)]
            parties.append({
                "name": name,
                "clean": _clean_counterparty_name(name),
                "type": ptype or "关联方",
                "aliases": aliases,
                "aliases_clean": [_clean_counterparty_name(a) for a in aliases],
            })
    return parties


# ---------------------------------------------------------------------------
# JSON 解析
# ---------------------------------------------------------------------------
def _parse_tables_json(tables_json: str) -> Tuple[List[str], List[List[str]], List[str], str]:
    """解析 LLM 传入 JSON，返回 (headers, rows, debit_credit_flags, error)。

    debit_credit_flags 与 rows 等长，可能为空字符串。
    """
    if not tables_json or not str(tables_json).strip():
        return [], [], [], "缺少 tables_json 参数"

    try:
        data = json.loads(tables_json) if isinstance(tables_json, str) else tables_json
    except json.JSONDecodeError as exc:
        return [], [], [], f"tables_json 不是合法 JSON：{exc}"

    headers = list(STANDARD_HEADERS)
    rows: List[List[str]] = []
    flags: List[str] = []

    def _append_row_from_mapping(mapping: Dict[str, Any]) -> None:
        row = [""] * len(STANDARD_HEADERS)
        for key, val in mapping.items():
            mapped = _map_header(key)
            if mapped and mapped != "借贷标志" and mapped in STANDARD_HEADERS:
                row[STANDARD_HEADERS.index(mapped)] = _cell_str(val)
        flag = ""
        for key, val in mapping.items():
            if _map_header(key) == "借贷标志":
                flag = _cell_str(val)
                break
        if any(row):
            rows.append(row)
            flags.append(flag)

    # 格式 1: {"headers": [...], "rows": [[...], ...]}
    if isinstance(data, dict):
        raw_headers = data.get("headers") or data.get("columns") or []
        raw_rows = data.get("rows") or data.get("data") or []
        if raw_headers and raw_rows and isinstance(raw_rows[0], (list, tuple)):
            col_map = []
            for h in raw_headers:
                mapped = _map_header(h) or (_cell_str(h) if _cell_str(h) in STANDARD_HEADERS else None)
                col_map.append(mapped)
            for raw_row in raw_rows:
                if not isinstance(raw_row, (list, tuple)):
                    continue
                row = [""] * len(STANDARD_HEADERS)
                flag = ""
                for idx, cell in enumerate(raw_row):
                    if idx >= len(col_map) or not col_map[idx]:
                        continue
                    if col_map[idx] == "借贷标志":
                        flag = _cell_str(cell)
                    elif col_map[idx] in STANDARD_HEADERS:
                        row[STANDARD_HEADERS.index(col_map[idx])] = _cell_str(cell)
                if any(row):
                    rows.append(row)
                    flags.append(flag)
            if rows:
                return headers, rows, flags, ""

        if raw_rows and isinstance(raw_rows[0], dict):
            for item in raw_rows:
                if isinstance(item, dict):
                    _append_row_from_mapping(item)
            if rows:
                return headers, rows, flags, ""

    if isinstance(data, list) and data and isinstance(data[0], dict):
        for item in data:
            if isinstance(item, dict):
                _append_row_from_mapping(item)
        if rows:
            return headers, rows, flags, ""
        return [], [], [], "JSON 数组中未找到有效交易明细行"

    if isinstance(data, list) and data and isinstance(data[0], (list, tuple)):
        header_row = [_cell_str(c) for c in data[0]]
        col_map = [_map_header(h) for h in header_row]
        if not any(col_map):
            col_map = [h if h in STANDARD_HEADERS else None for h in header_row]
        for raw_row in data[1:]:
            if not isinstance(raw_row, (list, tuple)):
                continue
            row = [""] * len(STANDARD_HEADERS)
            flag = ""
            for idx, cell in enumerate(raw_row):
                if idx >= len(col_map) or not col_map[idx]:
                    continue
                if col_map[idx] == "借贷标志":
                    flag = _cell_str(cell)
                elif col_map[idx] in STANDARD_HEADERS:
                    row[STANDARD_HEADERS.index(col_map[idx])] = _cell_str(cell)
            if any(row):
                rows.append(row)
                flags.append(flag)
        if rows:
            return headers, rows, flags, ""
        return [], [], [], "二维数组中未找到有效数据行"

    return [], [], [], "无法识别的 tables_json 结构，请使用 headers+rows 或对象数组格式"


# ---------------------------------------------------------------------------
# 3.1 标准化
# ---------------------------------------------------------------------------
def _split_income_expense(
    amount_raw: str,
    debit_credit: str = "",
    summary: str = "",
) -> Tuple[str, str]:
    """返回 (收入, 支出) 字符串。

    优先级：借贷标志 → 金额正负号 → 摘要关键词（无正负时）→ 默认正数作收入。
    摘要规则：货款/回款等→收入；材料款/采购款等→支出。
    """
    flag = debit_credit.strip()
    amt = _parse_amount(amount_raw)

    if flag:
        if any(x in flag for x in ("借", "D", "d", "支", "出")) and "贷" not in flag:
            if amt is None:
                return "", ""
            return "", _fmt_amount(abs(amt))
        if any(x in flag for x in ("贷", "C", "c", "收", "入")) and "借" not in flag:
            if amt is None:
                return "", ""
            return _fmt_amount(abs(amt)), ""

    if amt is None:
        return "", ""
    if amt < 0:
        return "", _fmt_amount(abs(amt))
    if amt > 0:
        # 无正负号且无借贷标志时，按摘要推断收支
        if not _amount_has_explicit_sign(amount_raw):
            direction = _infer_direction_from_summary(summary)
            if direction == "expense":
                return "", _fmt_amount(amt)
            if direction == "income":
                return _fmt_amount(amt), ""
        return _fmt_amount(amt), ""
    return "", ""


def step1_normalize(rows: List[List[str]], flags: Sequence[str]) -> List[Dict[str, str]]:
    """多银行流水格式标准化，输出 dict 行。"""
    out: List[Dict[str, str]] = []
    for i, row in enumerate(rows):
        rec = {h: (row[idx] if idx < len(row) else "") for idx, h in enumerate(STANDARD_HEADERS)}
        date_std = _parse_date(rec.get("交易日期", "")) or _cell_str(rec.get("交易日期", ""))
        time_std = _parse_time(rec.get("交易时间", "")) or _cell_str(rec.get("交易时间", ""))
        flag = flags[i] if i < len(flags) else ""
        summary = _cell_str(rec.get("摘要", ""))
        income, expense = _split_income_expense(rec.get("交易金额", ""), flag, summary)
        counterparty = _clean_counterparty_name(rec.get("对方户名", ""))
        # 若清洗后为空但原文有值，保留原文去空格
        if not counterparty and rec.get("对方户名"):
            counterparty = _cell_str(rec.get("对方户名", ""))

        out.append({
            "交易日期": date_std,
            "交易时间": time_std,
            "摘要": _cell_str(rec.get("摘要", "")),
            "收入": income,
            "支出": expense,
            "对方账号": _cell_str(rec.get("对方账号", "")),
            "对方户名": counterparty,
            "当前账户": _cell_str(rec.get("当前账户", "")),
            "当前户名": _cell_str(rec.get("当前户名", "")),
            "银行": _cell_str(rec.get("银行", "")),
        })
    return out


# ---------------------------------------------------------------------------
# 3.2 关联方识别
# ---------------------------------------------------------------------------
def _match_related_party(name: str, parties: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    clean = _clean_counterparty_name(name)
    if not clean:
        return None
    for p in parties:
        candidates = [p.get("clean") or "", p.get("name") or ""] + list(p.get("aliases_clean") or []) + list(p.get("aliases") or [])
        for c in candidates:
            c2 = _clean_counterparty_name(c) if c else ""
            if not c2:
                continue
            if clean == c2 or clean in c2 or c2 in clean:
                return p
    return None


def _is_controller_or_executive(ptype: str) -> bool:
    return any(k in (ptype or "") for k in ("实控人", "董监高", "董事", "监事", "高管", "实际控制人"))


def step2_related_party(
    records: List[Dict[str, str]],
    parties: List[Dict[str, Any]],
    config: Dict[str, Any],
) -> List[Dict[str, str]]:
    large_th = float(config["large_payment_threshold"])
    related_large = float(config["related_large_threshold"])
    out: List[Dict[str, str]] = []

    for rec in records:
        row = dict(rec)
        name = row.get("对方户名", "")
        matched = _match_related_party(name, parties) if parties else None
        income = _parse_amount(row.get("收入")) or 0.0
        expense = _parse_amount(row.get("支出")) or 0.0
        amount = max(income, expense)
        summary = row.get("摘要", "")

        if matched:
            row["是否关联方"] = "是"
            row["关联方类型"] = matched.get("type") or "关联方"
            risk = "中"
            reasons_high = False
            if _is_controller_or_executive(matched.get("type") or ""):
                reasons_high = True
            if amount >= related_large:
                reasons_high = True
            if any(k in summary for k in ("取现", "现金")):
                reasons_high = True
            # 关联方互转：对手也是关联方且本方户名像企业——已是关联方交易，大额或往来摘要抬高
            if any(k in summary for k in ("往来", "借款", "暂付", "互转")):
                reasons_high = True
            row["风险等级"] = "高" if reasons_high else "中"
        else:
            row["是否关联方"] = "否"
            row["关联方类型"] = ""
            # 非关联方但单笔超过大额阈值 → 中风险
            if amount >= large_th:
                row["风险等级"] = "中"
            else:
                row["风险等级"] = "低"
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# 3.3 异常交易筛查
# ---------------------------------------------------------------------------
def _is_night_time(time_str: str, night_start: str, night_end: str) -> bool:
    t = _parse_time(time_str)
    if not t:
        return False
    hh, mm = int(t[:2]), int(t[3:5])
    minutes = hh * 60 + mm

    def _to_min(s: str) -> int:
        parts = s.replace("：", ":").split(":")
        return int(parts[0]) * 60 + int(parts[1] if len(parts) > 1 else 0)

    start = _to_min(night_start)
    end = _to_min(night_end)
    # 23:00–05:00 跨日
    if start > end:
        return minutes >= start or minutes < end
    return start <= minutes < end


def _is_non_business(summary: str) -> bool:
    if any(k in summary for k in BUSINESS_KEYWORDS):
        return False
    if any(k in summary for k in NON_BUSINESS_KEYWORDS):
        return True
    return False


def step3_anomaly_screen(
    records: List[Dict[str, str]],
    config: Dict[str, Any],
) -> List[Dict[str, str]]:
    cash_th = float(config["cash_withdraw_threshold"])
    hf_amount = float(config["high_freq_amount"])
    hf_days = int(config["high_freq_days"])
    sensitive = list(config["sensitive_keywords"])
    night_start = str(config["night_start"])
    night_end = str(config["night_end"])

    # 预处理日期
    dated: List[Tuple[Optional[datetime], Dict[str, str]]] = []
    for rec in records:
        d = None
        if rec.get("交易日期"):
            try:
                d = datetime.strptime(rec["交易日期"], "%Y-%m-%d")
            except ValueError:
                d = None
        dated.append((d, rec))

    # R2：按对手方聚合非经营往来
    r2_hit_ids: Set[int] = set()
    by_party: Dict[str, List[Tuple[int, datetime, float]]] = defaultdict(list)
    for idx, (d, rec) in enumerate(dated):
        if d is None:
            continue
        if not _is_non_business(rec.get("摘要", "")):
            continue
        party = rec.get("对方户名") or ""
        if not party:
            continue
        amt = (_parse_amount(rec.get("收入")) or 0.0) + (_parse_amount(rec.get("支出")) or 0.0)
        if amt <= 0:
            continue
        by_party[party].append((idx, d, amt))

    for party, items in by_party.items():
        items.sort(key=lambda x: x[1])
        left = 0
        window_sum = 0.0
        for right, (idx, d, amt) in enumerate(items):
            window_sum += amt
            while left <= right and (d - items[left][1]).days > hf_days:
                window_sum -= items[left][2]
                left += 1
            if window_sum >= hf_amount:
                for j in range(left, right + 1):
                    r2_hit_ids.add(items[j][0])

    out: List[Dict[str, str]] = []
    for idx, (_d, rec) in enumerate(dated):
        row = dict(rec)
        hits: List[str] = []
        summary = row.get("摘要", "")
        expense = _parse_amount(row.get("支出")) or 0.0
        counterparty = row.get("对方户名", "")

        # R1 大额取现
        if expense >= cash_th and (
            "取现" in summary or counterparty in ("现金", "取现") or "现金" in counterparty
        ):
            hits.append(f"R1大额取现(支出≥{cash_th/10000:.0f}万)")

        # R2 高频往来
        if idx in r2_hit_ids:
            hits.append(f"R2高频往来({hf_days}天非经营累计≥{hf_amount/10000:.0f}万)")

        # R3 深夜交易
        if _is_night_time(row.get("交易时间", ""), night_start, night_end):
            hits.append(f"R3深夜交易({night_start}-{night_end})")

        # R4 敏感摘要
        for kw in sensitive:
            if kw and kw in summary:
                hits.append(f"R4敏感摘要({kw})")
                break

        if hits:
            row["是否异常"] = "是"
            row["命中原因"] = "；".join(hits)
            # 风险抬升：有异常且原为低 → 至少中；关联方异常 → 高
            if row.get("是否关联方") == "是" or row.get("风险等级") == "高":
                row["风险等级"] = "高"
            elif row.get("风险等级") == "低":
                row["风险等级"] = "中"
        else:
            row["是否异常"] = "否"
            row["命中原因"] = ""
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# 3.4 资金闭环核查
# ---------------------------------------------------------------------------
def _infer_issuer_names(records: List[Dict[str, str]], explicit: Optional[List[str]] = None) -> Set[str]:
    if explicit:
        return {_clean_counterparty_name(x) or x for x in explicit if x}
    counts: Dict[str, int] = defaultdict(int)
    for r in records:
        name = _cell_str(r.get("当前户名", ""))
        if name:
            counts[name] += 1
    if not counts:
        return set()
    # 取出现次数最多的本方户名作为发行人
    top = sorted(counts.items(), key=lambda x: -x[1])[0][0]
    return {top}


def step4_fund_loop(
    records: List[Dict[str, str]],
    parties: List[Dict[str, Any]],
    config: Dict[str, Any],
    issuer_names: Optional[List[str]] = None,
) -> List[Dict[str, str]]:
    large_th = float(config["large_payment_threshold"])
    window = int(config["trace_window_days"])
    issuers = _infer_issuer_names(records, issuer_names)
    related_names = set()
    for p in parties:
        related_names.add(p.get("clean") or p.get("name") or "")
        for a in (p.get("aliases_clean") or []) + (p.get("aliases") or []):
            related_names.add(_clean_counterparty_name(a) or a)
    related_names = {x for x in related_names if x}

    # 索引：按账户/户名
    dated_recs: List[Tuple[Optional[datetime], Dict[str, str]]] = []
    for rec in records:
        d = None
        try:
            if rec.get("交易日期"):
                d = datetime.strptime(rec["交易日期"], "%Y-%m-%d")
        except ValueError:
            d = None
        dated_recs.append((d, rec))

    chains: List[Dict[str, str]] = []

    for i, (d0, rec0) in enumerate(dated_recs):
        if d0 is None:
            continue
        # 发行人单笔支付超过阈值
        issuer_ok = False
        cur_name = _cell_str(rec0.get("当前户名", ""))
        if not issuers or cur_name in issuers:
            issuer_ok = True
        if not issuer_ok:
            continue
        expense = _parse_amount(rec0.get("支出")) or 0.0
        if expense < large_th:
            continue

        payee_name = rec0.get("对方户名", "")
        payee_acct = rec0.get("对方账号", "")
        if not payee_name and not payee_acct:
            continue

        # 追踪收款方在窗口内的转出
        window_end = d0 + timedelta(days=window)
        for j, (d1, rec1) in enumerate(dated_recs):
            if j == i or d1 is None:
                continue
            if d1 < d0 or d1 > window_end:
                continue
            # 收款方账户作为当前账户转出
            is_payee_out = False
            if payee_acct and rec1.get("当前账户") and payee_acct == rec1.get("当前账户"):
                is_payee_out = True
            if payee_name and rec1.get("当前户名") and (
                _clean_counterparty_name(rec1.get("当前户名", "")) == _clean_counterparty_name(payee_name)
                or payee_name == rec1.get("当前户名")
            ):
                is_payee_out = True
            # 也可能：同一对手在发行人流水里体现为对方收入后再支出——用对方户名匹配转出对手的下一跳
            # 简化：看 rec1 是否为该对手对外支付（当前户名=对手 或 无法跨户时，看是否有从对手账户出去）
            out_amt = _parse_amount(rec1.get("支出")) or 0.0
            if not is_payee_out or out_amt <= 0:
                continue

            end_name = rec1.get("对方户名", "")
            end_acct = rec1.get("对方账号", "")
            end_clean = _clean_counterparty_name(end_name)
            chain_type = "去向追踪"
            clue = "大额支付后收款方有资金转出"

            # 回流到发行人或关联方
            back_to_issuer = False
            if end_name and end_name in issuers:
                back_to_issuer = True
            if end_clean and end_clean in {_clean_counterparty_name(x) for x in issuers}:
                back_to_issuer = True
            if cur_name and end_name and end_name == cur_name:
                back_to_issuer = True
            if rec0.get("当前账户") and end_acct and end_acct == rec0.get("当前账户"):
                back_to_issuer = True

            back_to_related = bool(end_clean and end_clean in related_names) or bool(
                end_name and _match_related_party(end_name, parties)
            )

            if back_to_issuer:
                chain_type = "资金回流"
                clue = "资金回流至发行人账户，存在体外循环/回流信号"
            elif back_to_related:
                chain_type = "体外循环线索"
                clue = "资金流向关联方账户，存在体外循环线索"

            chains.append({
                "起点账户": rec0.get("当前账户", ""),
                "起点户名": cur_name,
                "起点银行": rec0.get("银行", ""),
                "支付日期": rec0.get("交易日期", ""),
                "支付金额": _fmt_amount(expense),
                "路径对手": payee_name,
                "路径摘要": rec0.get("摘要", ""),
                "终点账户": end_acct,
                "终点户名": end_name,
                "终点日期": rec1.get("交易日期", ""),
                "终点金额": _fmt_amount(out_amt),
                "链条类型": chain_type,
                "线索说明": clue,
            })

        # 若窗口内无转出，仍记录大额支付待追问
        if not any(
            c["支付日期"] == rec0.get("交易日期", "")
            and c["支付金额"] == _fmt_amount(expense)
            and c["路径对手"] == payee_name
            for c in chains
        ):
            chains.append({
                "起点账户": rec0.get("当前账户", ""),
                "起点户名": cur_name,
                "起点银行": rec0.get("银行", ""),
                "支付日期": rec0.get("交易日期", ""),
                "支付金额": _fmt_amount(expense),
                "路径对手": payee_name,
                "路径摘要": rec0.get("摘要", ""),
                "终点账户": "",
                "终点户名": "",
                "终点日期": "",
                "终点金额": "",
                "链条类型": "大额支付待追踪",
                "线索说明": f"单笔支付≥{large_th/10000:.0f}万，窗口内未匹配到收款方转出",
            })

    return chains


def _count_rule_hits(anomaly_rows: List[Dict[str, str]], rule_prefix: str) -> int:
    return sum(
        1
        for r in anomaly_rows
        if r.get("是否异常") == "是" and rule_prefix in (r.get("命中原因") or "")
    )


def _build_report_stats(
    normalized: List[Dict[str, str]],
    related_rows: List[Dict[str, str]],
    anomaly_rows: List[Dict[str, str]],
    chains: List[Dict[str, str]],
) -> Dict[str, Any]:
    """汇总核查成功时返回给 LLM 的 data 字段（不得由 LLM 编造）。"""
    banks = sorted({_cell_str(r.get("银行", "")) for r in normalized if _cell_str(r.get("银行", ""))})
    total_count = len(normalized)

    related_hits = [r for r in related_rows if r.get("是否关联方") == "是"]
    related_count = len(related_hits)
    related_high = sum(1 for r in related_hits if r.get("风险等级") == "高")

    party_counts: Dict[str, int] = defaultdict(int)
    for r in related_hits:
        name = _cell_str(r.get("对方户名", "")) or "（未知名）"
        party_counts[name] += 1
    top_related = [
        f"{name}（{cnt}笔）"
        for name, cnt in sorted(party_counts.items(), key=lambda x: (-x[1], x[0]))[:5]
    ]

    anomaly_count = sum(1 for r in anomaly_rows if r.get("是否异常") == "是")
    r1_count = _count_rule_hits(anomaly_rows, "R1")
    r2_count = _count_rule_hits(anomaly_rows, "R2")
    r3_count = _count_rule_hits(anomaly_rows, "R3")
    r4_count = _count_rule_hits(anomaly_rows, "R4")

    suspicious = [
        c for c in chains if c.get("链条类型") in ("资金回流", "体外循环线索", "大额支付待追踪")
    ]
    loop_count = len(suspicious)
    loop_details: List[str] = []
    for i, c in enumerate(suspicious, start=1):
        ctype = c.get("链条类型", "")
        start = c.get("起点户名") or c.get("起点账户") or "发行人"
        mid = c.get("路径对手") or "（对手未知）"
        end = c.get("终点户名") or c.get("终点账户") or "（去向未匹配）"
        pay_date = c.get("支付日期", "")
        pay_amt = c.get("支付金额", "")
        clue = c.get("线索说明", "")
        loop_details.append(
            f"{i}. [{ctype}] {pay_date} 支付{pay_amt}：{start} → {mid} → {end}"
            + (f"；{clue}" if clue else "")
        )

    follow_up: List[str] = []
    if related_count:
        follow_up.append(
            f"请核实 {related_count} 笔关联方往来的商业合理性、定价依据与审批流程"
            + (f"（含高风险 {related_high} 笔）" if related_high else "")
        )
    if top_related:
        follow_up.append(f"重点访谈主要关联方：{'、'.join(top_related)}")
    if r1_count:
        follow_up.append(f"请说明 R1 大额取现 {r1_count} 笔的资金用途、领取人与去向凭证")
    if r2_count:
        follow_up.append(f"请说明 R2 高频非经营往来 {r2_count} 笔的业务背景与资金闭环安排")
    if r3_count:
        follow_up.append(f"请说明 R3 深夜交易 {r3_count} 笔的发生原因与操作权限控制")
    if r4_count:
        follow_up.append(f"请对 R4 敏感摘要 {r4_count} 笔逐笔说明借款/往来/暂付等实质")
    loop_suspicious = [c for c in suspicious if c.get("链条类型") in ("资金回流", "体外循环线索")]
    if loop_suspicious:
        follow_up.append(
            f"请说明 {len(loop_suspicious)} 条回流/体外循环线索的业务实质与资金最终用途"
        )
    pending = [c for c in suspicious if c.get("链条类型") == "大额支付待追踪"]
    if pending:
        follow_up.append(f"请补充 {len(pending)} 笔大额支付收款方后续资金去向材料")
    if not follow_up:
        follow_up.append("未发现显著异常信号，建议抽样复核大额及敏感摘要交易")
    follow_up.append("AI 仅做雷达不做判决书，不做最终定性；以上为访谈追问方向，非审计结论")

    return {
        "bank_count": len(banks),
        "total_count": total_count,
        "related_count": related_count,
        "related_high": related_high,
        "top_related": top_related,
        "anomaly_count": anomaly_count,
        "r1_count": r1_count,
        "r2_count": r2_count,
        "r3_count": r3_count,
        "r4_count": r4_count,
        "loop_count": loop_count,
        "loop_details": loop_details,
        "follow_up": follow_up,
        "row_count": total_count,
        "chain_count": loop_count,
        "sheets": ["标准化流水", "关联方打标", "异常筛查", "资金闭环", "核查结论摘要"],
    }


def build_summary(stats: Dict[str, Any]) -> List[List[str]]:
    """将 report stats 写入 Excel「核查结论摘要」sheet。"""
    top_related = stats.get("top_related") or []
    loop_details = stats.get("loop_details") or []
    follow_up = stats.get("follow_up") or []
    lines = [
        ["指标", "数值"],
        ["核查银行数", str(stats.get("bank_count", 0))],
        ["流水总笔数", str(stats.get("total_count", 0))],
        ["关联方往来笔数", str(stats.get("related_count", 0))],
        ["高风险关联方笔数", str(stats.get("related_high", 0))],
        ["主要关联方", "；".join(top_related) if top_related else "无"],
        ["异常交易笔数", str(stats.get("anomaly_count", 0))],
        ["R1大额取现", str(stats.get("r1_count", 0))],
        ["R2高频往来", str(stats.get("r2_count", 0))],
        ["R3深夜交易", str(stats.get("r3_count", 0))],
        ["R4敏感摘要", str(stats.get("r4_count", 0))],
        ["可疑资金链条数", str(stats.get("loop_count", 0))],
        ["资金链条明细", " | ".join(loop_details) if loop_details else "无"],
        ["追问方向", "；".join(follow_up) if follow_up else "无"],
        ["说明", "AI 仅做雷达不做判决书，不做最终定性"],
    ]
    return lines


# ---------------------------------------------------------------------------
# Excel 输出
# ---------------------------------------------------------------------------
def _auto_width(ws, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    if not get_column_letter:
        return
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            if col_idx - 1 < len(row):
                max_len = max(max_len, len(str(row[col_idx - 1] or "")))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 48)


def _style_header(ws) -> None:
    if not Font:
        return
    fill = PatternFill("solid", fgColor=FILL_HEADER) if PatternFill else None
    for cell in ws[1]:
        cell.font = Font(bold=True)
        if fill:
            cell.fill = fill
        if Alignment:
            cell.alignment = Alignment(wrap_text=True, vertical="center")


def _row_fill(ws, row_idx: int, color: str, col_count: int) -> None:
    if not PatternFill:
        return
    fill = PatternFill("solid", fgColor=color)
    for c in range(1, col_count + 1):
        ws.cell(row=row_idx, column=c).fill = fill


def _dicts_to_rows(headers: Sequence[str], records: List[Dict[str, str]]) -> List[List[str]]:
    return [[_cell_str(r.get(h, "")) for h in headers] for r in records]


def _build_excel(
    normalized: List[Dict[str, str]],
    related_rows: List[Dict[str, str]],
    anomaly_rows: List[Dict[str, str]],
    chains: List[Dict[str, str]],
    summary_table: List[List[str]],
) -> bytes:
    wb = Workbook()

    # Sheet1 标准化流水
    ws1 = wb.active
    ws1.title = "标准化流水"
    ws1.append(list(NORMALIZED_HEADERS))
    _style_header(ws1)
    norm_rows = _dicts_to_rows(NORMALIZED_HEADERS, normalized)
    for r in norm_rows:
        ws1.append(r)
    _auto_width(ws1, NORMALIZED_HEADERS, norm_rows)

    # Sheet2 关联方打标
    ws2 = wb.create_sheet("关联方打标")
    ws2.append(list(RELATED_HEADERS))
    _style_header(ws2)
    rel_rows = _dicts_to_rows(RELATED_HEADERS, related_rows)
    for i, r in enumerate(rel_rows, start=2):
        ws2.append(r)
        rec = related_rows[i - 2]
        if rec.get("是否关联方") == "是":
            color = FILL_HIGH if rec.get("风险等级") == "高" else FILL_RELATED
            _row_fill(ws2, i, color, len(RELATED_HEADERS))
        elif rec.get("风险等级") == "中":
            _row_fill(ws2, i, FILL_MEDIUM, len(RELATED_HEADERS))
    _auto_width(ws2, RELATED_HEADERS, rel_rows)

    # Sheet3 异常筛查（仅异常行 + 全量也可；这里输出全量便于核对，异常行高亮）
    ws3 = wb.create_sheet("异常筛查")
    ws3.append(list(ANOMALY_HEADERS))
    _style_header(ws3)
    ano_rows = _dicts_to_rows(ANOMALY_HEADERS, anomaly_rows)
    for i, r in enumerate(ano_rows, start=2):
        ws3.append(r)
        rec = anomaly_rows[i - 2]
        if rec.get("是否异常") == "是":
            color = FILL_HIGH if rec.get("风险等级") == "高" else FILL_MEDIUM
            _row_fill(ws3, i, color, len(ANOMALY_HEADERS))
        elif rec.get("是否关联方") == "是":
            _row_fill(ws3, i, FILL_RELATED, len(ANOMALY_HEADERS))
    _auto_width(ws3, ANOMALY_HEADERS, ano_rows)

    # Sheet4 资金闭环
    ws4 = wb.create_sheet("资金闭环")
    ws4.append(list(CHAIN_HEADERS))
    _style_header(ws4)
    chain_rows = _dicts_to_rows(CHAIN_HEADERS, chains)
    for i, r in enumerate(chain_rows, start=2):
        ws4.append(r)
        ctype = chains[i - 2].get("链条类型", "")
        if ctype in ("资金回流", "体外循环线索"):
            _row_fill(ws4, i, FILL_HIGH, len(CHAIN_HEADERS))
        elif ctype == "大额支付待追踪":
            _row_fill(ws4, i, FILL_MEDIUM, len(CHAIN_HEADERS))
    _auto_width(ws4, CHAIN_HEADERS, chain_rows)

    # Sheet5 核查结论摘要
    ws5 = wb.create_sheet("核查结论摘要")
    for row in summary_table:
        ws5.append(row)
    _style_header(ws5)
    _auto_width(ws5, summary_table[0], summary_table[1:])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run_pipeline(
    tables_json: Any,
    related_parties: Any = None,
    config: Any = None,
    issuer_names: Any = None,
) -> Tuple[Optional[bytes], Dict[str, Any], str]:
    """执行完整四步处理，返回 (xlsx_bytes, stats, error)。"""
    if not isinstance(tables_json, str):
        tables_json = json.dumps(tables_json, ensure_ascii=False)

    _headers, rows, flags, err = _parse_tables_json(tables_json)
    if err:
        return None, {}, err

    cfg = _merge_config(config)
    parties = _parse_related_parties(related_parties)

    issuers: Optional[List[str]] = None
    if issuer_names:
        if isinstance(issuer_names, str):
            try:
                parsed = json.loads(issuer_names)
                issuers = parsed if isinstance(parsed, list) else [issuer_names]
            except json.JSONDecodeError:
                issuers = [x.strip() for x in re.split(r"[,，、]", issuer_names) if x.strip()]
        elif isinstance(issuer_names, list):
            issuers = [str(x) for x in issuer_names]

    normalized = step1_normalize(rows, flags)
    related_rows = step2_related_party(normalized, parties, cfg)
    anomaly_rows = step3_anomaly_screen(related_rows, cfg)
    chains = step4_fund_loop(anomaly_rows, parties, cfg, issuers)
    stats = _build_report_stats(normalized, related_rows, anomaly_rows, chains)
    summary_table = build_summary(stats)

    try:
        xlsx_bytes = _build_excel(normalized, related_rows, anomaly_rows, chains, summary_table)
    except Exception as e:  # noqa: BLE001
        logger.error("生成 Excel 失败: %s", e, exc_info=True)
        return None, {}, f"生成 Excel 失败：{e}"

    stats["summary"] = {row[0]: row[1] for row in summary_table[1:]}
    return xlsx_bytes, stats, ""


def main(input_data: Dict[str, Any], context: Dict[str, Any]) -> Dict[str, Any]:
    """解析 JSON → 四步核查处理 → 生成多 sheet Excel → 写盘并直接返回给 LLM。"""
    missing = _check_deps()
    if missing:
        msg = f"服务端缺少依赖库：{', '.join(missing)}，请管理员执行 pip install {' '.join(missing)}"
        return {"success": False, "message": msg}

    tables_json = input_data.get("tables_json") or ""
    file_name = input_data.get("file_name") or "银行流水核查结果.xlsx"
    related_parties = (
        input_data.get("related_parties")
        or input_data.get("关联方清单")
        or input_data.get("related_party_list")
    )
    config = input_data.get("config") or input_data.get("阈值配置")
    issuer_names = input_data.get("issuer_names") or input_data.get("发行人") or input_data.get("issuer")

    logger.info(
        "[json_to_excel] 收到参数: file_name=%s, related_parties=%s, tables_json_len=%s",
        file_name,
        bool(related_parties),
        len(tables_json) if isinstance(tables_json, str) else "obj",
    )

    if file_name.lower().endswith(".json"):
        file_name = file_name[:-5] + ".xlsx"
    elif not file_name.lower().endswith(".xlsx"):
        file_name = f"{file_name}.xlsx"

    xlsx_bytes, stats, err = run_pipeline(tables_json, related_parties, config, issuer_names)
    if err:
        return {"success": False, "message": err}

    if not xlsx_bytes.startswith(b"PK"):
        logger.error("[json_to_excel] 生成的内容不是有效 xlsx，前4字节=%s", xlsx_bytes[:4])
        return {"success": False, "message": "生成 Excel 失败：文件内容异常"}

    # 透传上游可选告警字段（仅存在时返回，禁止编造）
    optional_alerts: Dict[str, Any] = {}
    for key in ("pdf_convert_failures", "file_errors"):
        val = input_data.get(key)
        if val:
            if isinstance(val, str):
                try:
                    val = json.loads(val)
                except json.JSONDecodeError:
                    val = [val]
            if isinstance(val, list) and val:
                optional_alerts[key] = [str(x) for x in val]

    out_path = _resolve_output_path(file_name, context or {})
    try:
        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        with open(out_path, "wb") as fh:
            fh.write(xlsx_bytes)
    except Exception as e:  # noqa: BLE001
        logger.error("[json_to_excel] 写盘失败: %s", e, exc_info=True)
        return {"success": False, "message": f"Excel 已生成但写入失败：{e}"}

    file_size = len(xlsx_bytes)

    data = {
        "bank_count": int(stats.get("bank_count", 0)),
        "total_count": int(stats.get("total_count", 0)),
        "related_count": int(stats.get("related_count", 0)),
        "related_high": int(stats.get("related_high", 0)),
        "top_related": list(stats.get("top_related") or []),
        "anomaly_count": int(stats.get("anomaly_count", 0)),
        "r1_count": int(stats.get("r1_count", 0)),
        "r2_count": int(stats.get("r2_count", 0)),
        "r3_count": int(stats.get("r3_count", 0)),
        "r4_count": int(stats.get("r4_count", 0)),
        "loop_count": int(stats.get("loop_count", 0)),
        "loop_details": list(stats.get("loop_details") or []),
        "follow_up": list(stats.get("follow_up") or []),
        "file_name": file_name,
        "file_path": out_path,
        "file_size": file_size,
        "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "sheets": list(stats.get("sheets") or []),
    }
    data.update(optional_alerts)

    logger.info(
        "核查 Excel 完成: total=%s related=%s anomaly=%s loop=%s file_path=%s size=%s",
        stats.get("total_count"),
        stats.get("related_count"),
        stats.get("anomaly_count"),
        stats.get("loop_count"),
        out_path,
        file_size,
    )
    return {
        "success": True,
        "message": (
            f"银行流水核查完成（银行{stats.get('bank_count')}家，流水{stats.get('total_count')}笔，"
            f"关联方{stats.get('related_count')}笔，异常{stats.get('anomaly_count')}笔，"
            f"可疑链条{stats.get('loop_count')}条）。"
            f"Excel 已直接写盘输出给 LLM：file_path={out_path}；file_name={file_name}；"
            f"file_size={file_size}。未上传平台，无 download_url。"
            f"最终回复必须按 reference.md 规则 1：先输出一至五节核查内容，再交付 data.file_path 对应的 Excel；"
            f"字段只能取自 data，禁止编造；禁止自行拼接 /api/.../download 链接。"
            f"AI 仅做雷达不做判决书，不做最终定性。"
        ),
        "data": data,
    }


# ---------------------------------------------------------------------------
# 本地测试（平台不会走这里）
# ---------------------------------------------------------------------------
def run_local_test(
    tables_json: Any,
    output_xlsx: Optional[str] = None,
    file_name: str = "银行流水核查结果.xlsx",
    related_parties: Any = None,
    config: Any = None,
) -> Dict[str, Any]:
    missing = _check_deps()
    if missing:
        raise RuntimeError(
            f"缺少依赖库：{', '.join(missing)}，请执行 pip install {' '.join(missing)}"
        )

    xlsx_bytes, stats, err = run_pipeline(tables_json, related_parties, config)
    if err:
        return {"success": False, "message": err}

    if not file_name.lower().endswith(".xlsx"):
        file_name = f"{file_name}.xlsx"
    out_path = output_xlsx or file_name
    with open(out_path, "wb") as fh:
        fh.write(xlsx_bytes)

    return {
        "success": True,
        "message": f"已生成核查 Excel：{out_path}",
        "output_xlsx": out_path,
        "file_name": file_name,
        "file_size": len(xlsx_bytes),
        **stats,
    }


if __name__ == "__main__":
    import os
    import sys

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")

    SAMPLE_TABLES = {
        "headers": [
            "交易日期", "交易时间", "摘要", "交易金额",
            "对方账号", "对方户名", "当前账户", "当前户名", "银行",
        ],
        "rows": [
            ["2026/4/3", "09:15:00", "转支", "-510.00", "", "", "95599818******1111", "发行人A", "农业银行"],
            ["2026-04-05", "14:20:11", "货款", "3300.00", "", "客户甲有限公司", "95599818******1111", "发行人A", "农业银行"],
            ["20260408", "11:05:33", "转存", "6800.00", "", "", "95599818******1111", "发行人A", "农业银行"],
            ["2026-04-02", "16:40:00", "现存", "6500.00", "6228480168470007", "", "95599802******6666", "发行人A", "农业银行"],
            ["2026-04-09", "10:12:45", "往来款", "-2000000.00", "621661******7412", "张栋（本部）", "95599802******6666", "发行人A", "农业银行"],
            ["2026-04-12", "23:40:00", "借款", "-1200000.00", "622848******0266", "辰源贸易有限公司", "95599802******6666", "发行人A", "农业银行"],
            ["2026-04-13", "01:05:00", "往来", "-800000.00", "621700******5561", "李敏-分行", "95599802******6666", "发行人A", "农业银行"],
            ["2026-04-14", "10:00:00", "转账", "-1500000.00", "622200******8899", "中间方B", "95599802******6666", "发行人A", "农业银行"],
            ["2026-04-15", "11:00:00", "转账", "-1400000.00", "95599802******6666", "发行人A", "622200******8899", "中间方B", "工商银行"],
            ["2026-04-18", "15:55:01", "取现", "-600000.00", "", "现金", "95599802******6666", "发行人A", "农业银行"],
        ],
    }
    SAMPLE_PARTIES = [
        {"name": "张栋", "type": "实控人", "aliases": ["张*栋"]},
        {"name": "辰源贸易", "type": "关联企业", "aliases": ["辰源贸易有限公司"]},
        {"name": "李敏", "type": "董监高"},
    ]

    if len(sys.argv) >= 2:
        json_path = sys.argv[1]
        out_path = sys.argv[2] if len(sys.argv) > 2 else None
        with open(json_path, "r", encoding="utf-8") as fh:
            payload = fh.read()
        result = run_local_test(payload, output_xlsx=out_path, related_parties=SAMPLE_PARTIES)
    else:
        scripts_dir = os.path.dirname(os.path.abspath(__file__))
        out_path = os.path.join(scripts_dir, "银行流水核查_local_test.xlsx")
        result = run_local_test(SAMPLE_TABLES, output_xlsx=out_path, related_parties=SAMPLE_PARTIES)

    if not result.get("success"):
        print(f"生成失败: {result.get('message')}")
        sys.exit(1)

    print(f"银行数: {result.get('bank_count')}  总笔数: {result.get('total_count')}")
    print(f"关联方: {result.get('related_count')}(高{result.get('related_high')})  异常: {result.get('anomaly_count')}")
    print(f"R1-R4: {result.get('r1_count')}/{result.get('r2_count')}/{result.get('r3_count')}/{result.get('r4_count')}")
    print(f"链条: {result.get('loop_count')}  top_related: {result.get('top_related')}")
    print(f"大小: {result['file_size']} bytes")
    print(f"Excel 已写入: {result['output_xlsx']}")
