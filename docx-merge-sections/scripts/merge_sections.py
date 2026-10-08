# -*- coding: utf-8 -*-
"""
以主干文档为骨架，把源文档按【索引区间】切成的块交错插入指定位置。

用法:
    python merge_sections.py <master.docx> <source.docx> <plan.json> <output.docx>

plan.json 格式（ops 必须按 index 从大到小排列，即"从后往前"插入）：
{
  "chunks": [
    {"no":"01","name":"数据整合","start":656,"end":664,"index":"END"},
    {"no":"02","name":"历史数据处理","start":646,"end":655,"index":1501},
    {"no":"03","name":"某章节","start":100,"end":200,"index":900,
     "restyle":[["4","2"]]}
  ],
  "renames": [["历史数据","西点工作底稿数据迁移"]]
}

- index 为【主干文档原始】的 body child 索引；"END" 表示追加到 sectPr 之前
- restyle: [[from_styleId, to_styleId], ...]，用于层级对齐（如 H2->H3）
- renames: [[原标题, 新标题], ...]，精确全匹配且仅作用于 H2

关键约束：Composer.insert 内部 index+=1，连续插同一锚点时后插的排前面，
因此 ops 必须按 index 从大到小排列才能得到正序结果。
"""
import io
import sys
import os
import json

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docxcompose.composer import Composer

LVL = {'3': 'H1', '4': 'H2', '2': 'H3', '5': 'H4', '6': 'H5',
       '7': 'H6', '8': 'H7', '9': 'H8', '23': 'TITLE'}


def lvl_of(p):
    try:
        sid = p.style.style_id
    except Exception:
        sid = None
    if sid in LVL:
        return LVL[sid]
    pPr = p._p.find(qn('w:pPr'))
    if pPr is not None:
        ps = pPr.find(qn('w:pStyle'))
        if ps is not None and ps.get(qn('w:val')) in LVL:
            return LVL[ps.get(qn('w:val'))]
    return None


def restyle(doc, pairs):
    n = 0
    for frm, to in pairs:
        for p in doc.paragraphs:
            pPr = p._p.find(qn('w:pPr'))
            if pPr is None:
                continue
            ps = pPr.find(qn('w:pStyle'))
            if ps is not None and ps.get(qn('w:val')) == frm:
                ps.set(qn('w:val'), to)
                n += 1
    return n


def make_chunk(src, out, start, end, restyles=None):
    doc = Document(src)
    body = doc.element.body
    children = list(body)          # 先取快照
    for i in range(len(children) - 1, -1, -1):
        el = children[i]
        if el.tag.endswith('}sectPr'):
            continue               # 保留 sectPr
        if not (start <= i <= end):
            body.remove(el)
    n = restyle(doc, restyles) if restyles else 0
    doc.save(out)
    return n


def set_para_text(p, text):
    """保留首个 run 的格式，仅替换文字"""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''


def run(master_path, src_path, plan_path, out_path, workdir=None):
    plan = json.load(open(plan_path, encoding='utf-8'))
    workdir = workdir or os.path.join(os.path.dirname(out_path), '_merge_tmp')
    os.makedirs(workdir, exist_ok=True)

    master = Document(master_path)
    composer = Composer(master)
    n0 = len(master.element.body)
    print(f'主干: {master_path}  {n0} body children')

    total_added = 0
    for c in plan['chunks']:
        chunk_path = os.path.join(workdir, f"chunk_{c['no']}.docx")
        n_re = make_chunk(src_path, chunk_path, c['start'], c['end'], c.get('restyle'))
        idx = len(master.element.body) - 1 if c['index'] == 'END' else c['index']
        before = len(master.element.body)
        composer.insert(idx, Document(chunk_path))
        after = len(master.element.body)
        total_added += after - before
        extra = f'  层级调整 {n_re} 处' if n_re else ''
        print(f"  插入 chunk_{c['no']} @ {idx:<5d} 《{c['name']}》 "
              f"{before} -> {after} (+{after - before}){extra}")

    master.save(out_path)
    print(f'\n已保存 {out_path}')
    print(f'body children: {n0} -> {len(master.element.body)} (新增 {total_added})')

    # 后处理：改名
    if plan.get('renames'):
        doc = Document(out_path)
        hit = []
        for p in doc.paragraphs:
            t = p.text.strip()
            for old, new in plan['renames']:
                if t == old and lvl_of(p) == 'H2':
                    set_para_text(p, new)
                    hit.append(old)
        doc.save(out_path)
        print(f'改名完成: {hit}')
        miss = [o for o, _ in plan['renames'] if o not in hit]
        if miss:
            print(f'⚠️ 未匹配: {miss}')

    verify(out_path, n0 + total_added)


def verify(path, expect_body):
    print('\n=== 验收 ===')
    doc = Document(path)
    body = doc.element.body
    kids = list(body)
    tbls = len(body.findall(qn('w:tbl')))
    imgs = len(list(body.iter(qn('a:blip'))))
    words = len(''.join(t.text or '' for t in body.iter(qn('w:t'))))
    last = kids[-1].tag.split('}')[-1]

    ok = True
    for name, act, exp in [('body children', len(kids), expect_body),
                           ('末尾元素', last, 'sectPr')]:
        good = act == exp
        ok &= good
        print(f'  {"✅" if good else "❌"} {name} = {act} (期望 {exp})')
    print(f'  表格 = {tbls}   图片引用 = {imgs}   字数 = {words}')

    # 图片引用是否全部可解析
    import zipfile
    import posixpath
    from lxml import etree
    R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        rels = etree.fromstring(z.read('word/_rels/document.xml.rels'))
        rid2t = {r.get('Id'): r.get('Target') for r in rels}
        bad = 0
        for blip in body.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            rid = blip.get(f'{{{R}}}embed')
            if not rid:
                continue
            t = rid2t.get(rid)
            if t is None or posixpath.normpath(posixpath.join('word', t)) not in names:
                bad += 1
    print(f'  {"✅" if bad == 0 else "❌"} 图片引用失效 {bad} 个')
    ok &= (bad == 0)

    print('\n=== 章节快照 ===')
    for p in doc.paragraphs:
        lv = lvl_of(p)
        if lv in ('H1', 'H2', 'H3') and p.text.strip():
            print(f"{'':{0 if lv=='H1' else 4 if lv=='H2' else 8}}{lv}  {p.text.strip()[:56]}")

    print('\n' + ('✅ 验收通过' if ok else '❌ 存在未通过项'))


if __name__ == '__main__':
    if len(sys.argv) != 5:
        print(__doc__)
        sys.exit(1)
    run(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
